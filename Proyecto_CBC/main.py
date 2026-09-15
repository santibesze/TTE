

import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook



# Archivo que contiene los datos del transformador.
ARCHIVO_JSON = Path("entrada.json")

# Base de datos de los Conmutadores Bajo Carga.
ARCHIVO_CBC = Path("CBC (1).xlsx")

# Archivo Excel que generará el programa.
ARCHIVO_PMA = Path("PMA.xlsx")


def convertir_numero(valor):
    """
    Convierte un dato a número.

    El Excel puede contener:
        250
        "250"
        "250,5"
        "---"
        una celda vacía

    Si el valor no representa un número, devuelve None.
    """

    # Si la celda está vacía.
    if valor is None:
        return None

    # Si el dato es texto.
    if isinstance(valor, str):

        valor = valor.strip()

        # "---" significa que el dato no aplica.
        if valor == "" or valor == "---":
            return None

        # Permite números con coma decimal.
        valor = valor.replace(",", ".")

    try:
        return float(valor)

    except (ValueError, TypeError):
        return None


# ------------------------------------------------------------

def normalizar_conexion(grupo):
    """
    Convierte el grupo de conexión del transformador
    a la nomenclatura utilizada por la base de datos.

    Ejemplos:

        YN      -> Y
        yn/d11  -> Y
        D       -> D
        Δ       -> D
    """

    grupo = str(grupo).upper()

    # Reemplazamos el símbolo delta por D.
    grupo = grupo.replace("Δ", "D")

    if grupo.startswith("Y"):
        return "Y"

    if grupo.startswith("D"):
        return "D"

    raise ValueError(
        f"No se pudo determinar la conexión del transformador: {grupo}"
    )


# ------------------------------------------------------------

def todas_las_columnas_cumplen(fila, columnas, valor_requerido):
    """
    Comprueba un requisito contra varias columnas del CBC.

    Por ejemplo:

        a -> columnas V y W

    Si una celda contiene "---", significa que ese parámetro
    no aplica y se ignora.

    Si existe al menos un valor numérico, todos los valores
    aplicables deben ser mayores o iguales al valor requerido.
    """

    valores_validos = []

    for columna in columnas:

        valor = convertir_numero(fila[columna])

        # Ignoramos "---" y celdas vacías.
        if valor is not None:
            valores_validos.append(valor)

    # Si ninguna columna tiene información, no podemos
    # verificar el requisito.
    if len(valores_validos) == 0:
        return False

    # Todos los valores aplicables deben cumplir.
    for valor in valores_validos:

        if valor < valor_requerido:
            return False

    return True



def cargar_datos_transformador():
    """
    Abre entrada.json y devuelve los datos como un diccionario.
    """

    if not ARCHIVO_JSON.exists():

        raise FileNotFoundError(
            f"No se encontró el archivo {ARCHIVO_JSON}"
        )

    with open(
        ARCHIVO_JSON,
        "r",
        encoding="utf-8"
    ) as archivo:

        datos = json.load(archivo)

    return datos




def calcular_parametros(datos):
    """
    Realiza los cálculos necesarios para seleccionar el CBC.

    Parámetros obtenidos:

        E
        UF
        a
        b
        f
        V_step
        I_max
        P_step

    En esta versión se utiliza la configuración Y/estrella
    definida a partir del Excel original.

    NO se utilizan UWH1 ni UWH0.
    """

    # --------------------------------------------------------
    # DATOS DEL TRANSFORMADOR
    # --------------------------------------------------------

    # Potencia del transformador en MVA.
    S = convertir_numero(
        datos["potencia_prim"]
    )

    # Tensión primaria en kV.
    Un = convertir_numero(
        datos["tension_prim"]
    )

    # Cantidad de posiciones superiores.
    N1 = convertir_numero(
        datos["cbc_rango_pos"]
    )

    # Cantidad de posiciones inferiores.
    N2 = convertir_numero(
        datos["cbc_rango_neg"]
    )

    # Paso porcentual del CBC.
    e = convertir_numero(
        datos["cbc_paso_porc"]
    )

    # BIL de fase.
    UPH1 = convertir_numero(
        datos["bil_prim_fase"]
    )


    E = (
        ((N1 * e + N2 * e) * 100)
        /
        (
            200
            +
            (N1 * e - N2 * e)
        )
    )

    UF = (
        2 * E / (100 - E)
    ) + 0.1



    a = UF * UPH1

    b = UF * UPH1

    f = UF * UPH1


    V_step = (
        (Un * 1000)
        /
        math.sqrt(3)
    ) * (
        e / 100
    )



    I_max = (
        S * 1000
    ) / (
        math.sqrt(3)
        * Un
        * (
            1
            -
            N2 * e / 100
        )
    )



    P_step = (
        V_step
        * I_max
        / 1000
    )


    # Devolvemos todos los resultados.
    return {

        "E": E,

        "UF": UF,

        "a": a,

        "b": b,

        "f": f,

        "V_step": V_step,

        "I_max": I_max,

        "P_step": P_step
    }


# ============================================================
# CARGAR BASE DE DATOS DE CBC
# ============================================================

def cargar_base_cbc():
    """
    Lee el archivo Excel de los CBC.

    La primera fila del archivo contiene los nombres
    de las columnas.
    """

    if not ARCHIVO_CBC.exists():

        raise FileNotFoundError(
            f"No se encontró el archivo: {ARCHIVO_CBC}"
        )

    # pandas permite trabajar fácilmente con las filas
    # y columnas del Excel.
    tabla = pd.read_excel(
        ARCHIVO_CBC,
        header=0
    )

    return tabla



def seleccionar_cbc(
    tabla,
    calculos,
    conexion
):
    """
    Recorre todos los CBC de la base de datos y determina
    cuáles cumplen con los resultados calculados.

    Criterios:

    1. Conexión.
    2. Corriente nominal.
    3. Tensión máxima por escalón.
    4. Capacidad de contactos.
    5. Parámetro a.
    6. Parámetro b.
    7. Parámetro f.
    """

    # --------------------------------------------------------
    # Recuperar los resultados calculados.
    # --------------------------------------------------------

    a = calculos["a"]

    b = calculos["b"]

    f = calculos["f"]

    V_step = calculos["V_step"]

    I_max = calculos["I_max"]

    P_step = calculos["P_step"]


    # Lista donde guardaremos los CBC que cumplen.
    candidatos = []


    # --------------------------------------------------------
    # RECORRER TODOS LOS CBC
    # --------------------------------------------------------

    for _, fila in tabla.iterrows():

        # ====================================================
        # 1. CONEXIÓN
        # ====================================================

        conexion_cbc = str(
            fila["Conexión"]
        ).strip().upper()

        # Por si el Excel utiliza "DELTA".
        if conexion_cbc == "DELTA":
            conexion_cbc = "D"

        # Si la conexión no coincide, descartamos el CBC.
        if conexion_cbc != conexion:
            continue


        # ====================================================
        # 2. CORRIENTE NOMINAL
        # ====================================================

        corriente_cbc = convertir_numero(
            fila["Corriente nominal [A]"]
        )

        if corriente_cbc is None:
            continue

        # La corriente nominal del CBC debe ser mayor
        # o igual a la corriente calculada.
        if corriente_cbc < I_max:
            continue


        # ====================================================
        # 3. TENSIÓN POR ESCALÓN
        # ====================================================

        tension_ok = True

        columnas_tension = [

            "Maxima Tensión por escalón 1. [V]",

            "Maxima Tensión por escalón 2. [V]"
        ]

        for columna in columnas_tension:

            valor = convertir_numero(
                fila[columna]
            )

            # Si la columna tiene un valor, debe cumplir.
            # Si contiene "---", se ignora.
            if valor is not None:

                if V_step > valor:

                    tension_ok = False

                    break

        if not tension_ok:
            continue


        # ====================================================
        # 4. CAPACIDAD DE CONTACTOS
        # ====================================================

        capacidad_contactos = convertir_numero(
            fila["Capacidad de contactos. [kVA]"]
        )

        if capacidad_contactos is None:
            continue

        if P_step > capacidad_contactos:
            continue


        # ====================================================
        # 5. PARÁMETRO "a"
        # ====================================================

        columnas_a = [

            "a 1,2/50us [kV]",

            "a 50Hz 1min [kV]"
        ]

        if not todas_las_columnas_cumplen(
            fila,
            columnas_a,
            a
        ):
            continue


        # ====================================================
        # 6. PARÁMETRO "b"
        # ====================================================

        if conexion == "Y":


            columnas_b = [

                "b 1,2/50us [kV]",

                "b 50Hz 1min [kV]"
            ]

        else:

            columnas_b = [

                "b1 1,2/50us [kV]",

                "b1 50Hz 1min [kV]",

                "b2 1,2/50us [kV]",

                "b2 50Hz 1min [kV]",

                "b3 1,2/50us [kV]",

                "b3 50Hz 1min [kV]"
            ]

        if not todas_las_columnas_cumplen(
            fila,
            columnas_b,
            b
        ):
            continue


        # ====================================================
        # 7. PARÁMETRO "f"
        # ====================================================

        if conexion == "Y":



            columnas_f = [

                "f 1,2/50us [kV]",

                "f 50Hz 1min [kV]"
            ]

        else:



            columnas_f = [

                "f_1 1,2/50us [kV]",

                "f_1 50Hz 1min [kV]",

                "f_2 1,2/50us [kV]",

                "f_2 50Hz 1min [kV]"
            ]

        if not todas_las_columnas_cumplen(
            fila,
            columnas_f,
            f
        ):
            continue


        # ====================================================
        # CBC VÁLIDO
        # ====================================================

        candidatos.append(fila)


    # --------------------------------------------------------
    # ORDENAR CANDIDATOS
    # --------------------------------------------------------
    #
    # Primero se ordenan por corriente nominal.
    #
    # De esta forma, si un CBC de 250 A cumple y uno de
    # 400 A también cumple, se prioriza el de 250 A.
    #
    # --------------------------------------------------------

    candidatos.sort(
        key=lambda fila:
        convertir_numero(
            fila["Corriente nominal [A]"]
        )
    )


    return candidatos


# ============================================================
# CREAR / ACTUALIZAR PMA.XLSX
# ============================================================

def guardar_resultado(
    candidato,
    conexion,
    tension
):
    """
    Crea PMA.xlsx si no existe.

    Dentro del archivo crea la hoja:

        ACCESORIOS

    y escribe:

        A1 -> Modelo seleccionado
        B1 -> Conexión
        C1 -> Tensión

        A2 -> Modelo
        B2 -> Y o D
        C2 -> Tensión del transformador
    """

    # --------------------------------------------------------
    # Abrir o crear el archivo.
    # --------------------------------------------------------

    if ARCHIVO_PMA.exists():

        libro = load_workbook(
            ARCHIVO_PMA
        )

    else:

        libro = Workbook()


    # --------------------------------------------------------
    # Obtener o crear la hoja ACCESORIOS.
    # --------------------------------------------------------

    if "ACCESORIOS" in libro.sheetnames:

        hoja = libro["ACCESORIOS"]

    else:

        hoja = libro.create_sheet(
            "ACCESORIOS"
        )


    # --------------------------------------------------------
    # Encabezados
    # --------------------------------------------------------

    hoja["A1"] = "Modelo seleccionado"

    hoja["B1"] = "Conexión"

    hoja["C1"] = "Tensión [kV]"


    # --------------------------------------------------------
    # Resultados
    # --------------------------------------------------------

    hoja["A2"] = candidato["Modelo"]

    hoja["B2"] = conexion

    hoja["C2"] = tension


    # --------------------------------------------------------
    # Guardar archivo
    # --------------------------------------------------------

    libro.save(
        ARCHIVO_PMA
    )


# ============================================================
# MOSTRAR RESULTADOS
# ============================================================

def mostrar_resultados(
    datos,
    calculos,
    candidatos,
    conexion
):
    """
    Muestra en la terminal los resultados obtenidos.
    """

    print()
    print("=" * 60)
    print("RESULTADOS DEL CÁLCULO")
    print("=" * 60)

    print(
        f"E              = {calculos['E']:.4f} %"
    )

    print(
        f"UF             = {calculos['UF']:.6f}"
    )

    print(
        f"a              = {calculos['a']:.4f} kV"
    )

    print(
        f"b              = {calculos['b']:.4f} kV"
    )

    print(
        f"f              = {calculos['f']:.4f} kV"
    )

    print(
        f"V por escalón  = {calculos['V_step']:.4f} V"
    )

    print(
        f"I máxima       = {calculos['I_max']:.4f} A"
    )

    print(
        f"P por escalón  = {calculos['P_step']:.4f} kVA"
    )

    print(
        f"Conexión       = {conexion}"
    )

    print(
        f"Tensión        = {datos['tension_prim']} kV"
    )


    # --------------------------------------------------------
    # Mostrar candidatos.
    # --------------------------------------------------------

    print()
    print("=" * 60)
    print("CBC QUE CUMPLEN LOS CRITERIOS")
    print("=" * 60)

    for candidato in candidatos:

        print(
            f"Modelo: {candidato['Modelo']} | "
            f"Marca: {candidato['Marca']} | "
            f"Corriente: {candidato['Corriente nominal [A]']} A"
        )


# ============================================================
# FUNCIÓN PRINCIPAL
# ============================================================

def main():
    """
    Función principal del programa.

    Ejecuta todas las etapas en orden.
    """

    try:

        # ----------------------------------------------------
        # 1. Leer JSON
        # ----------------------------------------------------

        datos = cargar_datos_transformador()


        # ----------------------------------------------------
        # 2. Determinar conexión
        # ----------------------------------------------------

        conexion = normalizar_conexion(
            datos["prim_grupo"]
        )


        # ----------------------------------------------------
        # 3. Calcular parámetros
        # ----------------------------------------------------

        calculos = calcular_parametros(
            datos
        )


        # ----------------------------------------------------
        # 4. Cargar base de CBC
        # ----------------------------------------------------

        tabla_cbc = cargar_base_cbc()


        # ----------------------------------------------------
        # 5. Buscar CBC compatibles
        # ----------------------------------------------------

        candidatos = seleccionar_cbc(
            tabla_cbc,
            calculos,
            conexion
        )


        # ----------------------------------------------------
        # 6. Mostrar resultados
        # ----------------------------------------------------

        mostrar_resultados(
            datos,
            calculos,
            candidatos,
            conexion
        )


        # ----------------------------------------------------
        # 7. Verificar que exista un candidato
        # ----------------------------------------------------

        if len(candidatos) == 0:

            print()
            print(
                "NO SE ENCONTRÓ NINGÚN CBC COMPATIBLE."
            )

            return


        # ----------------------------------------------------
        # 8. Seleccionar el primer candidato
        # ----------------------------------------------------

        seleccionado = candidatos[0]


        print()
        print("=" * 60)
        print("CBC SELECCIONADO")
        print("=" * 60)

        print(
            f"Modelo: {seleccionado['Modelo']}"
        )

        print(
            f"Marca: {seleccionado['Marca']}"
        )

        print(
            f"Conexión: {conexion}"
        )


        # ----------------------------------------------------
        # 9. Crear / actualizar PMA.xlsx
        # ----------------------------------------------------

        guardar_resultado(

            seleccionado,

            conexion,

            datos["tension_prim"]
        )


        # ----------------------------------------------------
        # 10. Confirmación final
        # ----------------------------------------------------

        print()
        print(
            "Resultado guardado correctamente."
        )

        print(
            f"Archivo generado: {ARCHIVO_PMA}"
        )


    except Exception as error:

        print()
        print("=" * 60)
        print("ERROR")
        print("=" * 60)

        print(error)


# ============================================================
# EJECUTAR EL PROGRAMA
# ============================================================

if __name__ == "__main__":

    main()